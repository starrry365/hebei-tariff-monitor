#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把某网快照导成 Excel（板块分开 + 城市专属单列），供人工查阅/对外交付。

为什么要一个脚本而不是随手写几行：
  · 快照名带当天日期，**不能硬编码**（写死一天，第二天就导出一份旧数据还不报错）；
  · 字段口径（`_attr` 板块 / `_firstLevelName` 一级 / `_allCity` 不限地市 /
    `_cityNames` 城市专属）是构建期写好的，别在导出侧再算一遍 —— 两处算法迟早漂。

用法：
  python probes/tools/export_unicom_xlsx.py                  # 联通 → 默认输出到工作区
  python probes/tools/export_unicom_xlsx.py --net move       # 移动（4 位地市码那套）
  python probes/tools/export_unicom_xlsx.py --out D:/x.xlsx  # 指定输出

依赖：openpyxl（只本地用，不进 CI —— CI 不装它）。
"""
import argparse
import gzip
import json
import os
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
SNAP = os.path.join(REPO, "cloud", "tariff", "snapshots")

PREFIX = {"unicom": "unicom_tariff_", "move": "hebei_tariff_",
          "telecom": "ct_tariff_", "cbn": "cbn_tariff_"}
NETNAME = {"unicom": "联通", "move": "移动", "telecom": "电信", "cbn": "广电"}
BOARD = {"1": "全国", "2": "河北"}

COLS = [
    ("板块", 8), ("一级栏目", 12), ("二级栏目", 16), ("类型", 14),
    ("资费名称", 46), ("月费(元)", 10), ("流量", 12), ("语音(分钟)", 10),
    ("短信(条)", 9), ("宽带", 10), ("有效期", 22), ("上线日", 11), ("下线日", 11),
    ("状态", 8), ("适用人群", 18), ("销售渠道", 18), ("地域归属", 16), ("资费编号", 14),
]


def latest_snapshot(prefix):
    """该网**最新**的快照路径（按文件名里的 8 位日期排，不按 mtime）。"""
    cand = []
    for nm in os.listdir(SNAP):
        if nm.startswith(prefix) and nm.endswith(".json.gz"):
            cand.append((nm[len(prefix):-len(".json.gz")], os.path.join(SNAP, nm)))
    return max(cand) if cand else (None, None)


def yn(v):
    v = str(v or "").strip()
    return "" if v in ("", "0", "无", "None") else v


def row_of(e, board_map):
    attr = str(e.get("_attr") or "")
    fl = str(e.get("_firstLevel") or "")
    cty = e.get("_cityNames") or []
    return [
        board_map.get(attr, attr),
        e.get("_firstLevelName") or "",
        e.get("_secondLevelName") or "",
        e.get("type2Name") or "",
        e.get("name") or "",
        e.get("fees") or "",
        ("%s%s" % (yn(e.get("data")), yn(e.get("dataUnit")))) if yn(e.get("data")) else "",
        yn(e.get("minute")), yn(e.get("sms")),
        yn(e.get("broadBand")) or yn(e.get("brandwidth")),
        e.get("validPeriod") or "",
        e.get("onlineDay") or "", e.get("offineDay") or "",
        "已下架" if fl == "99" else "在售",
        e.get("applicablePeople") or "", e.get("channel") or "",
        "不限地市" if str(e.get("_allCity") or "") == "1"
        else ("、".join(cty) if cty else "不限地市"),
        e.get("reportNo") or "",
    ]


def write_sheet(wb, title, rows, note):
    ws = wb.create_sheet(title)
    hd = PatternFill("solid", fgColor="1E5AA8")
    for c, (name, w) in enumerate(COLS, 1):
        cell = ws.cell(1, c, name)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = hd
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    off = PatternFill("solid", fgColor="FDECEC")
    for i, r in enumerate(rows, 2):
        for c, v in enumerate(r, 1):
            cell = ws.cell(i, c, v)
            cell.alignment = Alignment(vertical="center", wrap_text=(c == 5),
                                       horizontal="right" if c in (6, 8, 9) else "left")
            if r[13] == "已下架":
                cell.fill = off
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(COLS)), len(rows) + 1)
    if note:
        ws.cell(len(rows) + 3, 1, note).font = Font(italic=True, size=9, color="777777")
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--net", default="unicom", choices=list(PREFIX))
    ap.add_argument("--out", default="")
    a = ap.parse_args()

    day, snap = latest_snapshot(PREFIX[a.net])
    if not snap:
        print("!! 找不到 %s 的快照（snapshots/%s*.json.gz）" % (NETNAME[a.net], PREFIX[a.net]))
        return 2
    d = json.load(gzip.open(snap, "rt", encoding="utf-8"))
    ent = d.get("entries") or []
    if not ent:
        print("!! 快照里没有 entries：%s" % os.path.basename(snap))
        return 2

    # 板块：联通/广电用 1/2 表示全国/本省；移动只有河北（无板块概念）
    if a.net == "move":
        bmap = {"1": "河北", "2": "河北", "3": "河北", "4": "河北", "5": "河北", "99": "河北"}
    else:
        bmap = dict(BOARD)

    boards = {}
    for e in ent:
        boards.setdefault(bmap.get(str(e.get("_attr") or ""), "?"), []).append(e)

    out = a.out or os.path.join(
        r"D:/Work/WorkBuddy", "%s资费专区-%s-%s.xlsx"
        % (NETNAME[a.net], "+".join(boards), day))

    wb = Workbook()
    wb.remove(wb.active)
    info = wb.create_sheet("说明")
    rows_info = [
        ("%s资费专区 · 全量数据" % NETNAME[a.net], ""),
        ("", ""),
        ("来源", d.get("src", "")),
        ("省份", "%s（provinceId=%s）" % (d.get("provinceName", ""), d.get("province", ""))),
        ("采集时间", d.get("fetchedAt", "")),
        ("是否含下架", "是" if d.get("includeStopped") else "否"),
        ("采集口径", "12 个地市的并集（该网资费按地市不同；单城会静默漏条目）"),
        ("", ""),
        ("总条数", len(ent)),
    ]
    for k, v in boards.items():
        rows_info.append(("  %s" % k, len(v)))
    rows_info += [
        ("", ""),
        ("不限地市（全省通用）", sum(1 for e in ent if str(e.get("_allCity")) == "1")),
        ("城市专属", sum(1 for e in ent if e.get("_cityNames"))),
        ("可筛地市", "、".join(c.get("cityName", "") for c in (d.get("cities") or []))),
        ("", ""),
        ("快照文件", os.path.basename(snap)),
        ("核查栏目覆盖", "python probes/probe_unicom_axes.py"),
    ]
    for i, (k, v) in enumerate(rows_info, 1):
        info.cell(i, 1, k).font = Font(bold=(i == 1), size=11 if i == 1 else 10)
        info.cell(i, 2, v).alignment = Alignment(wrap_text=True, vertical="center")
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 86

    write_sheet(wb, "全部(%d)" % len(ent), [row_of(e, bmap) for e in ent],
                "全部 %d 条，含已下架。" % len(ent))
    for k, v in boards.items():
        write_sheet(wb, "%s(%d)" % (k, len(v)), [row_of(e, bmap) for e in v],
                    "%s %d 条。" % (k, len(v)))

    city = [e for e in ent if e.get("_cityNames")]
    if city:
        write_sheet(wb, "城市专属(%d)" % len(city), [row_of(e, bmap) for e in city],
                    "只在部分地市目录里出现的资费，共 %d 条。" % len(city))

    wb.save(out)
    print("已导出：%s（%.0f KB）" % (out, os.path.getsize(out) / 1024))
    print("sheet：%s" % [ws.title for ws in wb.worksheets])
    return 0


if __name__ == "__main__":
    sys.exit(main())
